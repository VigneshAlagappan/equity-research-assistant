"""sources/rbi_reference_rate.py — parsing the RBI Reference Rate (INR per
USD) column out of the "Other Macroeconomic Indicators" workbook's Daily
sheet. Tests build a small synthetic workbook with the same shape as the
real publication rather than depending on the real downloaded file."""

from __future__ import annotations

from pathlib import Path

import openpyxl
import pytest

from sources.rbi_reference_rate import looks_like_rbi_reference_rate_workbook, parse_rbi_reference_rate


def _make_workbook(path: Path) -> None:
    wb = openpyxl.Workbook()
    wb.remove(wb.active)

    daily = wb.create_sheet("Daily")
    daily.append([None] * 9)
    daily.append([
        None, "Reporting Date", "NSE S&P CNX NIFTY", "BSE BANKEX", "REPO RATE (OVERNIGHT)",
        "REVERSE REPO RATE (OVERNIGHT)", "DAILY CALL MONEY RATE -HIGH", "DAILY CALL MONEY RATE - LOW",
        "RBI'S REFERENCE RATE: INR PER USD",
    ])
    daily.append([None, None, "Index", "Index", "Per cent", "Per cent", "Per cent", "Per cent", "Rupees"])
    daily.append([None, "24-Sep-2026", None, None, None, None, None, None, 95.9099])
    daily.append([None, "23-Sep-2026", None, None, None, None, None, None, 95.731])
    daily.append([None, "13-Sep-2026", "wh", "wh", "wh", "wh", "wh", "wh", "wh"])  # holiday row
    daily.append([None, "05-Sep-2026", None, None, 5.25, 3.35, 6, 3.25, None])  # no rate published yet

    weekly = wb.create_sheet("Weekly")
    weekly.append([None, "Reporting Date", "FOREIGN EXCHANGE RESERVES"])

    monthly = wb.create_sheet("Monthly")
    monthly.append([None, "Reporting Date", "M3"])

    wb.save(path)


@pytest.fixture
def workbook_path(tmp_path: Path) -> Path:
    path = tmp_path / "Other Macroeconomic Indicators.xlsx"
    _make_workbook(path)
    return path


def test_looks_like_rbi_reference_rate_workbook_true_for_matching_shape(workbook_path: Path) -> None:
    assert looks_like_rbi_reference_rate_workbook(workbook_path) is True


def test_looks_like_rbi_reference_rate_workbook_false_for_other_files(tmp_path: Path) -> None:
    wb = openpyxl.Workbook()
    wb.active.title = "Sheet1"
    path = tmp_path / "not_it.xlsx"
    wb.save(path)
    assert looks_like_rbi_reference_rate_workbook(path) is False


def test_looks_like_rbi_reference_rate_workbook_false_for_rbi_indicator_workbook() -> None:
    """The "50 Macroeconomic Indicators" workbook (Weekly/Fortnightly/
    Monthly/Quarterly, no Daily sheet) must not be mistaken for this shape."""
    import tempfile

    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    wb.create_sheet("Weekly")
    wb.create_sheet("Fortnightly")
    wb.create_sheet("Monthly")
    wb.create_sheet("Quarterly")
    with tempfile.TemporaryDirectory() as tmp:
        path = Path(tmp) / "50 Macroeconomic Indicators.xlsx"
        wb.save(path)
        assert looks_like_rbi_reference_rate_workbook(path) is False


def test_parses_reference_rate_with_string_dates(workbook_path: Path) -> None:
    obs = parse_rbi_reference_rate(workbook_path)
    assert {o.period for o in obs} == {"2026-09-24", "2026-09-23"}
    assert all(o.series_key == "usd_inr_reference_rate" for o in obs)
    assert all(o.period_type == "dated" for o in obs)
    assert all(o.unit == "INR_PER_USD" for o in obs)
    assert all(o.source == "rbi" for o in obs)


def test_holiday_and_blank_rows_are_skipped(workbook_path: Path) -> None:
    obs = parse_rbi_reference_rate(workbook_path)
    periods = {o.period for o in obs}
    assert "2026-09-13" not in periods  # "wh" holiday placeholder
    assert "2026-09-05" not in periods  # blank reference-rate cell


def test_units_sub_row_is_not_mistaken_for_a_data_row(workbook_path: Path) -> None:
    obs = parse_rbi_reference_rate(workbook_path)
    assert all(o.value != 0 for o in obs)
    assert len(obs) == 2
