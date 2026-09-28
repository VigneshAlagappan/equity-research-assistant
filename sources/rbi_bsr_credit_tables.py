"""Parses RBI's Basic Statistical Return (BSR) outstanding-credit tables
(Handbook of Statistics table 1.x series) -- quarterly, wide-pivoted
exports distributed as standalone downloads: category rows (Population
Group / Occupation / Type of Account) x a repeating block of date columns,
each date itself spanning several metric sub-columns (No. of Accounts,
Credit Limit, Amount Outstanding, ...).

Not the same shape as sources/rbi_dbie_tables.py's single-table parser
(period x item matrix, one column per period) or sources/rbi_indicators.py
(period x indicator matrix, one row per period) -- here BOTH periods and
metrics repeat across columns, with categories down the rows. A single
generic parser handles all three known tables (1.1, 1.4, 1.5) since they
share this exact layout, just with a different date-block width (4 vs 3
metric columns per date) and category vocabulary.
"""

from __future__ import annotations

import calendar
import logging
import re
from datetime import datetime
from pathlib import Path

import openpyxl

from sources.macro import MacroNormalizedObservation

logger = logging.getLogger(__name__)

PARSER_VERSION = "rbi-bsr-credit-tables-v1-xlsx"

_TABLE_TAG_RE = re.compile(r"TABLE\s+NO\.?\s*(\d+[.\-]\d+)", re.IGNORECASE)
_MONTH_YEAR_RE = re.compile(r"^([A-Za-z]{3})-(\d{4})$")


def _unit_for_metric(metric_label: str) -> str:
    label = metric_label.lower()
    if "account" in label or "offices" in label:
        return "NUMBER"
    return "INR_CRORE"  # "Credit Limit" / "Amount Outstanding"


def _slugify(label: str) -> str:
    label = label.replace("–", "-").replace("’", "")
    label = re.sub(r"[^A-Za-z0-9]+", "_", label)
    return re.sub(r"_+", "_", label).strip("_").lower()


def _find_table_tag(ws) -> str | None:
    for row in ws.iter_rows(min_row=1, max_row=4, values_only=True):
        for value in row:
            if isinstance(value, str):
                match = _TABLE_TAG_RE.search(value)
                if match:
                    return "t" + match.group(1).replace("-", "_").replace(".", "_")
    return None


def _find_dates_row(ws) -> int | None:
    """The row headed "As on End" (table 1.1, 1.4) or, in table 1.5's
    layout, headed by the category label instead ("Type of Account") --
    detected by content (>=2 parseable dates from column C onward), not by
    a fixed label string, since that label isn't consistent across tables."""
    for row_num, row in enumerate(ws.iter_rows(min_row=1, max_row=10), start=1):
        date_cells = [c for c in row[2:] if _period_from_date_cell(c.value) is not None]
        if len(date_cells) >= 2:
            return row_num
    return None


def _period_from_date_cell(value: object) -> str | None:
    if isinstance(value, datetime):
        return value.strftime("%Y-%m-%d")
    if isinstance(value, str):
        match = _MONTH_YEAR_RE.match(value.strip())
        if match:
            month_abbr, year = match.groups()
            month = datetime.strptime(month_abbr, "%b").month
            last_day = calendar.monthrange(int(year), month)[1]
            return f"{year}-{month:02d}-{last_day:02d}"
    return None


def looks_like_rbi_bsr_credit_table(file_path: Path) -> bool:
    try:
        wb = openpyxl.load_workbook(file_path, read_only=True)
    except Exception:
        return False
    ws = wb[wb.sheetnames[0]]
    return _find_table_tag(ws) is not None and _find_dates_row(ws) is not None


def parse_rbi_bsr_credit_table(file_path: Path) -> list[MacroNormalizedObservation]:
    wb = openpyxl.load_workbook(file_path, data_only=True)
    ws = wb[wb.sheetnames[0]]

    table_tag = _find_table_tag(ws)
    dates_row_num = _find_dates_row(ws)
    if table_tag is None or dates_row_num is None:
        raise ValueError(f"{file_path}: not a recognized BSR credit-table shape")

    dates_row = [c.value for c in ws[dates_row_num]]
    metrics_row = [c.value for c in ws[dates_row_num + 1]]

    # date blocks: columns (0-indexed into the row tuples) whose dates_row
    # cell is a non-None date/month-year string. The block width is the
    # gap between the first two such columns.
    date_cols = [i for i, v in enumerate(dates_row) if i >= 2 and _period_from_date_cell(v) is not None]
    if len(date_cols) < 2:
        raise ValueError(f"{file_path}: could not find at least 2 date columns in the 'As on End' row")
    stride = date_cols[1] - date_cols[0]

    data_start_row = dates_row_num + 2
    # A row of plain column numbers (1, 2, 3, ...) sometimes follows the
    # metrics row -- not data, skip it rather than hardcoding an offset.
    first_data_row = [c.value for c in ws[data_start_row]]
    if first_data_row[1] is None and first_data_row[2] == 1:
        data_start_row += 1

    observations: list[MacroNormalizedObservation] = []
    for row in ws.iter_rows(min_row=data_start_row, max_row=ws.max_row, values_only=True):
        category = row[1]
        if not isinstance(category, str) or not category.strip():
            continue
        category_key = _slugify(category)
        for date_col in date_cols:
            period = _period_from_date_cell(dates_row[date_col])
            if period is None:
                continue
            for offset in range(stride):
                col = date_col + offset
                if col >= len(row) or col >= len(metrics_row):
                    continue
                metric_label = metrics_row[col]
                value = row[col]
                if not isinstance(metric_label, str) or not isinstance(value, (int, float)) or isinstance(value, bool):
                    continue
                observations.append(MacroNormalizedObservation(
                    series_key=f"{table_tag}_{category_key}_{_slugify(metric_label)}",
                    period_type="dated", period=period, value=float(value), unit=_unit_for_metric(metric_label),
                    source="rbi", source_file=str(file_path), parser_version=PARSER_VERSION,
                ))
    return observations
