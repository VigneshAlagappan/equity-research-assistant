"""Parses RBI's "Other Macroeconomic Indicators" workbook's Daily sheet --
specifically RBI's own official Reference Rate (INR per USD), the one
INR/USD series actually sourced from RBI rather than the display-only
yfinance quote in web/fx_rate.py (that module's own docstring says it is
not authoritative for valuation math) -- this closes that gap in
macro_observations.

Only this one column is extracted, not the rest of the Daily sheet (Nifty/
Bankex index levels, call money rates) or this workbook's Weekly/Monthly
sheets (forex reserves, M1/M2/M3, RBI balance sheet detail, ...) -- those
overlap series already ingested from the "50 Macroeconomic Indicators"
workbook or are out of scope for now. A later pass through this module is
the natural place to add them, not a reason to import them speculatively
today.
"""

from __future__ import annotations

import logging
import re
from datetime import datetime
from pathlib import Path

import openpyxl

from sources.macro import MacroNormalizedObservation

logger = logging.getLogger(__name__)

PARSER_VERSION = "rbi-reference-rate-v1-xlsx"

SERIES_KEY = "usd_inr_reference_rate"
_DATE_RE = re.compile(r"^\d{1,2}-[A-Za-z]{3}-\d{4}$")


def _find_header(ws) -> tuple[int, int, int]:
    """(header_row, reporting_date_col, reference_rate_col). Searches the
    first few rows only -- same defensive-and-cheap shape check
    sources/rbi_indicators.py's _find_header_row uses."""
    for row_num, row in enumerate(ws.iter_rows(min_row=1, max_row=5), start=1):
        header_cells = {cell.value: cell.column for cell in row if isinstance(cell.value, str)}
        date_col = header_cells.get("Reporting Date")
        if date_col is None:
            continue
        rate_col = next(
            (
                col for value, col in header_cells.items()
                if "REFERENCE RATE" in value.upper() and "USD" in value.upper()
            ),
            None,
        )
        if rate_col is not None:
            return row_num, date_col, rate_col
    raise ValueError("could not find a 'Reporting Date' + \"...Reference Rate...USD...\" header row")


def _parse_date(raw: object) -> str | None:
    if isinstance(raw, datetime):
        return raw.strftime("%Y-%m-%d")
    if isinstance(raw, str) and _DATE_RE.match(raw.strip()):
        return datetime.strptime(raw.strip(), "%d-%b-%Y").strftime("%Y-%m-%d")
    return None


def looks_like_rbi_reference_rate_workbook(file_path: Path) -> bool:
    """Cheap, defensive shape check -- used to route between this module
    and sources/rbi_dbie_tables.py's single-table parser without guessing."""
    try:
        # Not read_only=True: this workbook's <dimension> XML tag undercounts
        # its own columns (confirmed against the real RBI export -- read-only
        # mode trusts that tag and silently truncates row iteration before
        # the Reference Rate column), unlike sources/rbi_indicators.py's
        # workbook where read_only is reliable.
        wb = openpyxl.load_workbook(file_path)
    except Exception:
        return False
    if "Daily" not in wb.sheetnames:
        return False
    try:
        _find_header(wb["Daily"])
    except ValueError:
        return False
    return True


def parse_rbi_reference_rate(file_path: Path) -> list[MacroNormalizedObservation]:
    """Extracts RBI's Reference Rate (INR per USD) from the Daily sheet.
    Raises ValueError if the workbook doesn't have the expected shape --
    same defensive check parse_rbi_indicator_workbook() makes."""
    wb = openpyxl.load_workbook(file_path, data_only=True)
    if "Daily" not in wb.sheetnames:
        raise ValueError(f"{file_path}: no 'Daily' sheet found")
    ws = wb["Daily"]
    header_row, date_col, rate_col = _find_header(ws)

    observations: list[MacroNormalizedObservation] = []
    skipped = 0
    # min_row=header_row + 1 (not +2): the row directly below the header is
    # a units sub-row ("Rupees", "Per cent", ...), not data -- it's dropped
    # for free by _parse_date() returning None for its blank Reporting Date
    # cell, same as any other malformed row, rather than being hardcoded as
    # a fixed offset that would break if a future export drops that row.
    for row in ws.iter_rows(min_row=header_row + 1, max_row=ws.max_row):
        period = _parse_date(row[date_col - 1].value)
        if period is None:
            skipped += 1
            continue
        value = row[rate_col - 1].value
        if not isinstance(value, (int, float)) or isinstance(value, bool):
            continue  # blank cells and the "wh" (holiday) placeholder land here too
        observations.append(MacroNormalizedObservation(
            series_key=SERIES_KEY, period_type="dated", period=period,
            value=float(value), unit="INR_PER_USD", source="rbi", source_file=str(file_path),
            parser_version=PARSER_VERSION,
        ))

    if skipped:
        logger.warning(
            "%s [Daily]: skipped %d row(s) with an unparseable Reporting Date", file_path, skipped
        )
    return observations
